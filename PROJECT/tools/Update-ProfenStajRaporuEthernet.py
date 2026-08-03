"""Add the Ethernet FFT and final delivery chapters to the internship report.

Run with the bundled workspace Python so python-docx and its dependencies are
reproducible. The update is idempotent: an older copy of the chapter is removed
before the current text is inserted immediately before Appendix A.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "PROJECT" / "reports" / "profen_staj_raporu.docx"
CHAPTER = "26. Sürekli Bare-Metal UDP FFT Aktarımı — 31 Temmuz 2026"
DELIVERY_CHAPTER = "27. Proje Teslim Kapanışı — 3 Ağustos 2026"


def remove_existing_chapter(document: Document) -> None:
    start = None
    end = None
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == CHAPTER:
            start = paragraph._p
        elif start is not None and paragraph.text.strip().startswith("Ek A"):
            end = paragraph._p
            break
    if start is None:
        return

    current = start
    while current is not None and current is not end:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def add_before(marker, text: str, style: str | None = None):
    paragraph = marker.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def remove_matching_paragraphs(document: Document, prefixes: tuple[str, ...]) -> None:
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith(prefixes):
            paragraph._p.getparent().remove(paragraph._p)


def enable_field_updates(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    document = Document(REPORT)
    remove_existing_chapter(document)
    remove_matching_paragraphs(document, ("26. Sürekli Bare-Metal UDP FFT", "27. Proje Teslim Kapanışı"))

    toc_marker = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Ekler A")
    )
    add_before(toc_marker, "26. Sürekli Bare-Metal UDP FFT Aktarımı", "TOC Entry")
    add_before(toc_marker, "27. Proje Teslim Kapanışı", "TOC Entry")

    marker = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Ek A")
    )

    add_before(marker, CHAPTER, "Heading 1")
    add_before(
        marker,
        "Bu ek geliştirme, kart üzerinde doğrulanmış 2048-bin PL FFT sonucunu "
        "PetaLinux eklemeden, AXU4EVB-P PS GEM3 arabirimi ve lwIP RAW UDP "
        "üzerinden Windows bilgisayara sürekli taşımak için hazırlanmıştır. "
        "Çalışan Golden, vitis2 ve önceki Candidate paketleri değiştirilmemiştir.",
    )

    add_before(marker, "26.1 Mimari ve veri sözleşmesi", "Heading 2")
    for item in (
        "Kart: 192.168.1.10/24; bilgisayar: 192.168.1.100/24; UDP hedef portu: 5005.",
        "Her FFT karesi 2048 adet unsigned 64-bit spektrum bininden oluşur.",
        "64 baytlık sürümlü başlık network byte order, spektrum payload'u little-endian kullanır.",
        "Bir kare, Ethernet MTU sınırını aşmayan 12 UDP datagramına bölünür: ilk 11 parça 176, son parça 112 bin taşır.",
        "Varsayılan yayın saniyede 4 kare ve süresizdir; --frames N yalnız kontrollü kayıt ve kabul testi içindir.",
    ):
        add_before(marker, item, "List Bullet")

    add_before(marker, "26.2 Tamamlanan uygulama", "Heading 2")
    for item in (
        "Transport bağımsız sürekli FFT servisi ve parçalama/protokol katmanı.",
        "ZynqMP GEM3 ile lwIP211 RAW/statik IPv4 taşıyıcısı.",
        "Mevcut Vitis projelerinden ayrılmış vitis_ethernet çalışma alanı ve Release ELF üretimi.",
        "Python datagram decoder, frame reassembly, canlı Matplotlib FFT görüntüleyici ve kayıt seçenekleri.",
        "Sıra, parça kapsamı, CRC, frame bütünlüğü ve zaman sınırı denetimli headless kabul/soak araçları.",
    ):
        add_before(marker, item, "List Bullet")

    add_before(marker, "26.3 Doğrulama sonucu ve dürüst sınır", "Heading 2")
    add_before(
        marker,
        "Host regresyonundaki 38 test geçmiştir. Vitis 2020.1 ile temiz ayrı "
        "çalışma alanında gerçek Release build'i tamamlanmış ve "
        "vitis_ethernet/workspace/ad9364app/Release/ad9364app.elf üretilmiştir. "
        "Buna karşılık bilgisayarda Ethernet portu ve USB-Ethernet adaptörü "
        "bulunmadığı için kart ping'i, gerçek UDP frame yakalama, 150-frame "
        "kabul, 600 saniyelik soak ve kablosuz CW ölçümü henüz yapılmamıştır. "
        "Dolayısıyla bu bölüm fiziksel Ethernet/RF PASS veya board_verified=true "
        "iddiası içermez; monopulse-live-ethernet-v1 Candidate henüz yayımlanmamıştır.",
    )

    add_before(marker, "26.4 Adaptör geldikten sonraki devir adımları", "Heading 2")
    for index, item in enumerate(
        (
            "Kartı JTAG, PS_UART ve Ethernet üzerinden bağla; PuTTY kaydını aç.",
            "USB-Ethernet adaptörünü 192.168.1.100/24 yap; gateway ve DNS alanlarını boş bırak.",
            "PROJECT/tools/Start-Axu4evAd9364EthernetDev.cmd ile Ethernet ELF'ini başlat.",
            "UART'ta PL_FFT_ETH: READY satırını gör ve ping 192.168.1.10 kapısını geç.",
            "Önce 150 ardışık frame headless kabulünü, sonra en az 600 saniye soak testini çalıştır.",
            "Ağ yolu kararlıysa SMW200A ile düşük seviyeden başlayan, mesafeli ve sınırlı kablosuz CW kabulünü uygula.",
            "Tüm kanıtlar PASS olduktan sonra monopulse-live-ethernet-v1 Candidate paketini oluştur ve hash/log/manifest ile arşivle.",
        ),
        start=1,
    ):
        add_before(marker, f"{index}. {item}")

    add_before(marker, DELIVERY_CHAPTER, "Heading 1")
    add_before(
        marker,
        "Staj kapanışında kaynak kodu, doğrulanmış Golden/Candidate paketleri, "
        "otomatik testler, sunum ve teknik raporlar tek bir teslim düzeninde "
        "toplanmıştır. Kök README proje haritasını; "
        "PROJECT/reports/TESLIM_NOTLARI_2026-08-03.md ise devralacak kişinin "
        "ilk gün sırasını, değişmez artefaktları ve açık kabul kapılarını tanımlar.",
    )

    add_before(marker, "27.1 Son yazılım doğrulaması", "Heading 2")
    for item in (
        "Ethernet decoder/reassembly/viewer regresyonu: 38 test PASS.",
        "Monopulse, FFT C-model ve UART model regresyonu: 61 test PASS.",
        "Bare-metal FFT lab-console regresyonu: 39 test PASS.",
        "Golden release, live-lab Candidate ve periyodik monitor paket bütünlüğü: PASS.",
        "Teslim RAR'ı temiz Git snapshotı, tam Git bundle ve SHA-256 manifestiyle üretilir.",
    ):
        add_before(marker, item, "List Bullet")

    add_before(marker, "27.2 Teslimde açık bırakılan işler", "Heading 2")
    add_before(
        marker,
        "USB-Ethernet adaptörü bulunmadığından fiziksel UDP kabulü ve kablosuz CW "
        "Ethernet testi tamamlanmamıştır. ZX80-DR230-S+ ile EV1HMC936ALP6 S0-S3 "
        "kontrol entegrasyonu da gelecekteki çalışmadır. Eski Workspace launcher, "
        "farklı tarihlerden kalan Vivado/Vitis kopyalarının birebir hash eşitliğini "
        "aradığı için teslim yolu değildir; günlük kullanımda immutable Golden ve "
        "Candidate launcher'ları esas alınır. Bu sınırlar ürün kabulü yapılmış gibi "
        "yorumlanmamalıdır.",
    )

    enable_field_updates(document)
    document.core_properties.modified = datetime.now()
    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
