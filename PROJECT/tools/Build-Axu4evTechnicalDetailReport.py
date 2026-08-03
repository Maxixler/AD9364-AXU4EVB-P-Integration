"""Build the AXU4EVB-P + AD9364 technical handoff report."""

from __future__ import annotations

import os
import subprocess
import tempfile
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "PROJECT" / "reports" / "profen_teknik_detay_raporu.docx"
NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "EEF1F4"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"


def git_text(*args: str) -> str:
    return subprocess.check_output(["git", "-C", str(ROOT), *args], text=True).strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Sayfa ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "İçindekiler alanını Word/LibreOffice içinde güncelleyin."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = text
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(8)
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
        if len(table.rows) % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        number = p.add_run(f"{index}. ")
        number.bold = True
        p.add_run(item)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.keep_together = True
    p.style = doc.styles["Normal"]
    set_cell = OxmlElement("w:shd")
    set_cell.set(qn("w:fill"), "F3F5F7")
    p._p.get_or_add_pPr().append(set_cell)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def add_status(doc: Document, title: str, text: str, color: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(f"{title}: ")
    r.bold = True
    p.add_run(text)
    doc.add_paragraph()


def draw_flow(path: Path, title: str, labels: list[str], colors: list[str] | None = None) -> None:
    colors = colors or ["#D9EAF7"] * len(labels)
    fig, ax = plt.subplots(figsize=(12, 2.1))
    ax.set_xlim(0, len(labels) * 2.0)
    ax.set_ylim(0, 2.0)
    ax.axis("off")
    for i, label in enumerate(labels):
        x = i * 2.0 + 0.15
        box = FancyBboxPatch(
            (x, 0.55), 1.55, 0.72,
            boxstyle="round,pad=0.08,rounding_size=0.08",
            linewidth=1.5, edgecolor="#2F75B5", facecolor=colors[i]
        )
        ax.add_patch(box)
        ax.text(x + 0.775, 0.91, label, ha="center", va="center", fontsize=8.5, weight="bold")
        if i < len(labels) - 1:
            ax.annotate("", xy=(x + 1.92, 0.91), xytext=(x + 1.58, 0.91),
                        arrowprops=dict(arrowstyle="->", color="#17365D", lw=1.6))
    ax.set_title(title, fontsize=12, color="#17365D", weight="bold", pad=4)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_chapter(doc: Document, number: int, title: str) -> None:
    heading = doc.add_heading(f"{number}. {title}", level=1)
    if number > 1:
        heading.paragraph_format.page_break_before = True


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in [
        ("Title", 26, NAVY), ("Heading 1", 17, NAVY),
        ("Heading 2", 13, BLUE), ("Heading 3", 10.5, NAVY)
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(2)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "PROFEN | AXU4EVB-P + AD9364 Teknik Devir Raporu"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(BLUE)
    add_page_number(section.footer.paragraphs[0])


def build_report() -> Path:
    branch = git_text("branch", "--show-current")
    commit = git_text("rev-parse", "HEAD")
    short_commit = commit[:12]
    doc = Document()
    configure_document(doc)

    with tempfile.TemporaryDirectory(prefix="axu4ev-tech-report-") as td:
        temp = Path(td)
        flow_system = temp / "system-flow.png"
        flow_fft = temp / "fft-flow.png"
        flow_jtag = temp / "jtag-flow.png"
        flow_udp = temp / "udp-flow.png"
        draw_flow(flow_system, "Uçtan Uca Sistem", ["RF / Anten", "AD9364", "ADI HDL", "PL FFT", "A53\nBare-metal", "UART / UDP", "Python"])
        draw_flow(flow_fft, "PL FFT Veri Yolu", ["I/Q Tap", "2048 Örnek\nCapture", "xFFT v9.1", "Magnitude²", "Peak", "Spectrum\nBRAM"])
        draw_flow(flow_jtag, "Doğrulanmış JTAG Başlatma Sırası", ["Connect", "PMU Gate", "PMUFW", "Bitstream", "A53 Reset", "PS/DDR Init", "DDR Test", "ELF", "Run"])
        draw_flow(flow_udp, "Sürekli Ethernet FFT Taşıması", ["Atomik Frame", "64 B Başlık", "12 UDP\nDatagram", "CRC / Sıra", "Reassembly", "Viewer /\nSoak"])

        # Cover
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(90)
        r = p.add_run("PROFEN")
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("AXU4EVB-P + AD9364\nFPGA FFT ve Monopulse Platformu")
        r.bold = True
        r.font.size = Pt(27)
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Teknik Tasarım, Uygulama ve Devir Raporu")
        r.italic = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        doc.add_paragraph()
        doc.add_picture(str(flow_system), width=Inches(6.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(45)
        p.add_run(f"Teslim tarihi: 3 Ağustos 2026\nDal: {branch}\nRapor üretim tabanı: {short_commit}")
        doc.add_page_break()

        doc.add_heading("Doküman Kontrolü", level=1)
        add_table(doc, ["Alan", "Değer"], [
            ["Doküman", "AXU4EVB-P + AD9364 Teknik Tasarım, Uygulama ve Devir Raporu"],
            ["Amaç", "Projeyi devralan mühendise uygulanmış mimariyi, çalışma adımlarını ve doğrulama sınırlarını aktarmak"],
            ["Kapsam", "Vivado, Vitis, JTAG bring-up, AD9364, PL FFT, UART/Python, UDP/Ethernet ve monopulse yol haritası"],
            ["Kaynak dal", branch], ["Rapor üretim tabanı", commit],
            ["Araç sürümü", "Vivado/Vitis 2020.1; PowerShell 5.1+; Python 3"],
            ["Durum", "Teslim sürümü — fiziksel Ethernet ve RF switch entegrasyonu açık"],
        ], [4.1, 12.7])
        doc.add_heading("Dokümanın kullanım sınırı", level=2)
        doc.add_paragraph(
            "Bu rapor yalnız teknik tasarımı ve uygulanan mühendislik adımlarını içerir. "
            "Staj sürecinin kurumsal/genel anlatımı profen_staj_raporu.docx dosyasındadır. "
            "Çelişki halinde kart şemaları, etkin Vivado projesi, derleme raporları, "
            "Candidate manifestleri ve kart kabul kayıtları önceliklidir."
        )
        doc.add_heading("İçindekiler", level=1)
        add_toc(doc.add_paragraph())

        add_chapter(doc, 1, "Teknik kapsam ve teslim kararı")
        doc.add_paragraph(
            "Proje, ALINX AXU4EVB-P kartındaki Zynq UltraScale+ MPSoC ile AD9364 tabanlı "
            "RF ön ucunu birleştirerek FPGA içinde 2048 noktalı spektrum üretmek, peak "
            "bulmak ve gelecekte dört analog Σ/Δ durumu üzerinden monopulse hata ölçümü "
            "oluşturmak üzere geliştirilmiştir. Gerçek zamanlı örnek işleme PL'de; RF "
            "konfigürasyonu, başlatma, kontrol ve sonuç aktarımı A53 bare-metal yazılımında yapılır."
        )
        add_status(doc, "Teslim sonucu", "Kart bring-up, AD9364 Rev 2 başlatma, tek-frame ve 100-frame PL FFT, UART/Python canlı laboratuvar hattı ve Ethernet yazılım/build zinciri doğrulandı.", GREEN)
        add_status(doc, "Açık fiziksel kapılar", "USB-Ethernet adaptörüyle gerçek UDP soak, kontrollü kablosuz CW Ethernet kabulü ve ZX80/HMC936 S0–S3 elektriksel entegrasyonu tamamlanmadı.", AMBER)
        add_table(doc, ["İş paketi", "Durum", "Bağlayıcı kanıt"], [
            ["Golden kart bring-up", "PASS", "PROJECT/releases/axu4ev-ad9364/2026-07-16-bringup-v1"],
            ["AD9364 Rev 2 + DMA", "PASS", "PS_UART başarı işaretçileri"],
            ["2048 nokta PL FFT — tek frame", "PASS", "fft-single-shot-v1"],
            ["2048 nokta PL FFT — 100 frame", "PASS", "fft-periodic-monitor-v1 + ham UART kanıtı"],
            ["Monopulse foundation", "PASS", "MONOPULSE_FOUNDATION: PASS"],
            ["UART/Python live-lab", "PASS", "monopulse-live-lab-v2 + test paketi"],
            ["Sürekli UDP FFT yazılım/build", "PASS", "38 host testi + ayrı Vitis Ethernet Release ELF"],
            ["Fiziksel Ethernet/UDP", "AÇIK", "Adaptör, ping, 150-frame ve 600 s soak bekleniyor"],
            ["ZX80/HMC936 S0–S3", "GELECEK", "GPIO/XDC/level-shifter ve RF kablo matrisi bekleniyor"],
        ], [5.2, 2.2, 9.2])

        add_chapter(doc, 2, "Donanım platformu ve bağlayıcı kısıtlar")
        add_table(doc, ["Bileşen", "Teknik rol / doğrulanan bilgi"], [
            ["AXU4EVB-P", "Zynq UltraScale+ MPSoC taşıyıcı kart; hedef parça xczu4ev-sfvc784-1-i"],
            ["AD9364/FMCOMMS kartı", "Tek fiziksel RX zinciri; RXB kullanılan RF girişi; LVDS dijital arayüz"],
            ["R&S SMW200A", "Düşük seviyeli, limitli CW RF kaynağı; kablosuz kabul için planlandı"],
            ["ZX80-DR230-S+", "DC–3 GHz SPDT RF switch; iki analog hata yolundan seçim"],
            ["EV1HMC936ALP6", "1.2–1.4 GHz IF faz kaydırıcı değerlendirme kartı; 6-bit kontrol; seviye uyumluluğu kapısı"],
            ["JTAG + PS_UART", "Programlama/başlatma ve 115200-8-N-1 çalışma kanıtı"],
            ["GEM3 Ethernet", "Bare-metal lwIP RAW UDP; kart 192.168.1.10, PC 192.168.1.100/24"],
        ], [4.1, 12.7])
        doc.add_heading("2.1 AD9364 mimari etkisi", level=2)
        add_bullets(doc, [
            "AD9364 tek RX içerdiğinden iki bağımsız anten kanalı aynı anda dijital olarak örneklenemez.",
            "Nihai monopulse yaklaşımı eşzamanlı iki dijital kanal yerine analog ön uçta sıralı S0–S3 durumları üretir.",
            "Bağlayıcı dört durum: Σ+ΔAZ, Σ−ΔAZ, Σ+ΔEL ve Σ−ΔEL.",
            "Frame'ler arası kompleks Δ/Σ faz oranı ilk ürün iddiası değildir; ilk ürün ortak peak çevresindeki güç tabanlı normalize AZ/EL oranıdır.",
            "Mutlak dBm ve açı sonucu, RF/kablo/phase-code kalibrasyonu tamamlanmadan raporlanmaz."
        ])
        doc.add_heading("2.2 RF güvenlik sınırı", level=2)
        doc.add_paragraph(
            "Kablosuz CW testinde jeneratör düşük çıkış seviyesinden başlatılmalı, cihazın seviye limiti "
            "etkin tutulmalı, anten mesafesi kontrollü artırılmalı ve AD9364 girişinde aşırı güç/clipping "
            "gözlenirse test derhal durdurulmalıdır. Bu rapor ölçülmemiş bir serbest-uzay kaybını güvenlik "
            "garantisi olarak kabul etmez."
        )

        add_chapter(doc, 3, "Kaynak ağacı ve değişmez teslimler")
        add_table(doc, ["Yol", "İçerik / kullanım"], [
            ["Related-Docs/hdl-main/projects/fmcomms2/axu4ev/fmcomms2_axu4ev.xpr", "Etkin Vivado 2020.1 projesi"],
            ["Related-Docs/hdl-main/library/axi_ad9364_fft_capture", "Özel capture/FFT/spectrum AXI IP kaynakları"],
            ["Vitis/workspace/ad9364app", "Golden/temel bare-metal uygulama kaynakları"],
            ["vitis2/workspace/ad9364app", "Güncel Candidate uygulama kaynakları"],
            ["vitis_ethernet/workspace/ad9364app", "Ayrı Ethernet geliştirme workspace'i"],
            ["PROJECT/releases", "Yerinde değiştirilmeyen Golden paketleri"],
            ["PROJECT/candidates", "Tek-frame, periyodik ve live-lab paketleri"],
            ["PROJECT/lab_console", "UART codec, seri istemci, model, viewer ve kayıt"],
            ["PROJECT/ethernet_fft", "UDP codec, frame reassembly, viewer ve soak araçları"],
            ["PROJECT/tests", "C/Python/PowerShell regresyon ve fiziksel kanıtlar"],
            ["PROJECT/reports", "Tasarım, plan, durum, sunum ve teslim raporları"],
        ], [7.4, 9.4])
        add_status(doc, "Değişmezlik kuralı", "Golden ve mevcut Candidate içerikleri yerinde düzenlenmez. Her yeni özellik yeni dal, yeni Candidate adı, yeni manifest/hash ve ayrı kabul kaydı üretir.", LIGHT_BLUE)
        doc.add_heading("3.1 Git ve yeniden üretilebilirlik", level=2)
        doc.add_paragraph(
            "Teslim RAR'ında izlenen kaynak snapshotına ek olarak GIT_GECMISI/FPGA-Profen.git.bundle "
            "bulunur. Bu bundle dal ve commit geçmişini taşır. Related-Docs/no-OS eski bir gitlink'tir; "
            "bağımsız checkout boş kalabilir. Kullanılan AD9361/AD9364 sürücüleri Vitis uygulama "
            "kaynaklarında vendored durumdadır."
        )
        add_code(doc, "git clone .\\GIT_GECMISI\\FPGA-Profen.git.bundle FPGA-Profen\ncd FPGA-Profen\ngit switch codex/axu4ev-monopulse-lab-console")

        add_chapter(doc, 4, "Vivado donanım tasarımı ve uygulanan düzeltmeler")
        doc.add_heading("4.1 Aktif tasarımın teyidi", level=2)
        add_numbered(doc, [
            "Etkin XPR, hedef FPGA parçası, system_top üst modülü ve system_constr.xdc dosyası belirlendi.",
            "AXU4EVB-P ve AD9364 şemaları üzerinden FMC/LVDS netleri üçlü eşleme yöntemiyle kontrol edildi: şema neti → XDC satırı → top-level port.",
            "ADI HDL veri yolunun mevcut RX DMAC yolunu beslediği ve özel FFT yolunun geri basınç oluşturmadan shadow tap kullanması gerektiği doğrulandı.",
            "Temiz synthesis, implementation, timing ve DRC koşuları ile bitstream/XSA üretildi."
        ])
        doc.add_heading("4.2 Uygulanan XDC düzeltmeleri", level=2)
        add_table(doc, ["Düzeltme", "Teknik gerekçe"], [
            ["RX_FRAME pinleri L7/L6", "AXU4EVB-P şeması ve FMC netiyle uyum"],
            ["TX_CLK/FB_CLK pinleri N7/N6", "Geri-besleme saat çiftinin doğru fiziksel eşlenmesi"],
            ["16 RX LVDS porta DIFF_TERM_ADV TERM_100", "Diferansiyel terminasyonun açık ve tutarlı olması"],
            ["SPI chip-select pull-up", "Başlatma sırasında kontrolsüz CS seçimini önleme"],
            ["rx_clk için 4.00 ns", "250 MHz giriş saatinin timing analizine bağlanması"],
        ], [6.0, 10.8])
        doc.add_heading("4.3 Özel PL FFT yolu", level=2)
        doc.add_picture(str(flow_fft), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_bullets(doc, [
            "Tap noktası util_ad9361_adc_fifo çıkışı / util_ad9361_adc_pack girişi arasındadır; ADI'nin mevcut CDC'si korunur.",
            "Mantıksal kompleks tuple (I0,Q0); 32-bit word {Q0[15:0], I0[15:0]} ve little-endian byte sırası I_lo, I_hi, Q_lo, Q_hi'dır.",
            "Capture, FFT TREADY'den bağımsız 2048 geçerli örneği RAM'e yazar; replay yalnız TVALID && TREADY ile ilerler.",
            "Mevcut pack + RX DMAC hattı FFT busy/backpressure nedeniyle durdurulamaz.",
            "Tek-frame kabulünden sonra atomik ping-pong spectrum bankası ve frame ownership kuralları kullanıldı."
        ])

        add_chapter(doc, 5, "FFT sayısal sözleşmesi")
        add_table(doc, ["Parametre", "Bağlayıcı değer"], [
            ["FFT çekirdeği", "Xilinx xFFT v9.1"],
            ["Dönüşüm boyutu", "2048 kompleks örnek"],
            ["Çıkış sırası", "Natural order"],
            ["Scaling schedule", "0x6AB"],
            ["Spectrum sözcüğü", "2048 × unsigned 64-bit magnitude-square"],
            ["Peak", "peak_bin, peak_power, diagnostik peak_re/peak_im"],
            ["Frame bütünlüğü", "frame_id, sample_count, raw-frame CRC-32, reset_epoch, active_bank"],
            ["Güç ekseni", "Relative dB, uncalibrated"],
        ], [5.5, 11.3])
        doc.add_heading("5.1 Frekans eşleme", level=2)
        doc.add_paragraph(
            "Python görünümü fftshift uygulanmış eksende çalışır. Bin aralığı Δf = Fs/N'dir; "
            "Fs=30.72 MSPS ve N=2048 için 15 kHz/bin elde edilir. RF ekseni, merkez LO ile "
            "baseband bin ofsetinin toplamıdır. DC guard ve arama bandı peak seçiminden önce uygulanabilir."
        )
        add_table(doc, ["Örnek RF profili", "Değer"], [
            ["LO", "2.400 GHz"], ["Sample rate", "30.72 MSPS"],
            ["RX bandwidth", "10 MHz"], ["RX gain", "71 dB manual"],
            ["TX attenuation", "89.75 dB / TX kullanılmıyor"], ["FDD", "RX sürekliliği için kullanılabilir"],
        ], [6.0, 10.8])
        doc.add_heading("5.2 Tek-frame ve sürekli frame davranışı", level=2)
        add_bullets(doc, [
            "Tek-frame modunda tetik, 2048 örnek capture, FFT, spectrum bank commit ve status okuma sıralıdır.",
            "Periyodik mod 100 ardışık frame'i sıra/CRC/peak alanlarıyla üretir ve tam 100 frame sonunda PASS verir.",
            "Live-lab modu host istekli çalışır; her CAPTURE_ONCE en fazla bir tam spectrum döndürür.",
            "Ethernet modu varsayılan olarak süresizdir; frame limiti yalnız açık parametreyle uygulanır."
        ])

        add_chapter(doc, 6, "Vitis platformu ve bare-metal yazılım")
        doc.add_heading("6.1 Platform onarımı", level=2)
        add_numbered(doc, [
            "Vivado'dan bitstream içeren güncel XSA üretildi ve Vitis hardware platform bu XSA ile yenilendi.",
            "Out-of-date platform durumu, PMU firmware hedefinin psu_pmu_0/mb-gcc olarak düzeltilmesi ve XPFM yeniden üretimiyle giderildi.",
            "FSBL OCM taşması bölüm bazlı derleme seçenekleriyle çözüldü; FSBL ELF üretildi.",
            "ad9364app eski sdr_wrapper bağlantılarından ayrılıp fmcomms2_axu4ev_hw / standalone_domain hedefine bağlandı.",
            "Release include/library yolları ve libm bağlantısı düzeltildi; AD9364 seçimi AD9361_DEVICE=0, AD9364_DEVICE=1 olarak doğrulandı.",
            "AArch64 Release ELF derlendi; IDE bitstream ve psu_init dosyaları Candidate ile eşlendi."
        ])
        doc.add_heading("6.2 Uygulama modları", level=2)
        add_table(doc, ["Compile-time mod", "Amaç / beklenen işaretçi"], [
            ["DMA_EXAMPLE", "Temel RX DMA kabulü; Done."],
            ["WIFI_RF_TEST", "2.4 GHz çevresel aktivite taraması; WIFI_RF_TEST: PASS"],
            ["MONOPULSE_FOUNDATION_TEST", "RF profil ve temel aritmetik; MONOPULSE_FOUNDATION: PASS"],
            ["PL_FFT_SINGLE_SHOT", "Tek atomik spectrum; PL_FFT: PASS frame=..."],
            ["PL_FFT_PERIODIC_MONITOR", "100 sıralı frame ve kapanış PASS"],
            ["PL_FFT_LAB_UART", "Binary UART istek/yanıt servisi"],
            ["PL_FFT_ETHERNET", "GEM3/lwIP üzerinden sürekli UDP spectrum"],
        ], [5.5, 11.3])

        add_chapter(doc, 7, "Doğrulanmış JTAG bring-up")
        doc.add_picture(str(flow_jtag), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_numbered(doc, [
            "Yerel hw_server'a bağlanılır ve hedef ağacı okunur.",
            "PSU hedefinde 0xFFCA0038 adresine 0x1FF yazılarak PMU JTAG güvenlik kapıları açılır.",
            "MicroBlaze PMU seçilir, zynqmp_pmufw/executable.elf yüklenir ve çalıştırılır.",
            "Doğrulanmış system_top.bit ile PL programlanır.",
            "Cortex-A53 #0 resetlenir ve register'lar temizlenir.",
            "psu_init.tcl kaynak edilir; psu_init, PS–PL isolation removal ve reset config uygulanır.",
            "DDR 0x00100000 adresinde A5A5A5A5 yaz/oku kapısıyla doğrulanır.",
            "Seçilen Golden veya Candidate ad9364app.elf A53'e yüklenir.",
            "A53 çalıştırılır; sonuç PS_UART üzerinden doğrulanır."
        ])
        add_code(doc, ".\\PROJECT\\tools\\Start-Axu4evAd9364Golden.cmd\n.\\PROJECT\\tools\\Start-Axu4evAd9364Candidate.cmd\n.\\PROJECT\\tools\\Start-Axu4evAd9364EthernetDev.cmd")
        doc.add_heading("7.1 Başlıca arıza belirtileri ve kök neden", level=2)
        add_table(doc, ["Belirti", "Kök neden / doğru işlem"], [
            ["Instruction transfer timeout", "Kart/target durumu kilitli; tam güç döngüsü, tek hw_server, PMU gate ve sıralı init"],
            ["Memory write error at 0x0", "DDR/PS init tamamlanmadan A53 ELF indirme girişimi"],
            ["PMU-FW is not running", "PMUFW launcher'a eklenmeli; PMU hedefinde dow + con"],
            ["Unsupported PRODUCT_ID 0x00/0xFF", "PS–PL isolation/reset veya SPI/RF kart başlatma sırası eksik"],
            ["PL_FFT FAIL CAPTURE -6", "Kart/RF yolunu yeniden bağlayıp tam Candidate başlatma; kalıcıysa capture/status incelemesi"],
            ["Platform Out-of-date", "Hardware platform regenerate/build; yalnız IDE index problemiyle karıştırılmamalı"],
        ], [6.0, 10.8])

        add_chapter(doc, 8, "Kart kabulü ve RF doğrulama adımları")
        doc.add_heading("8.1 Temel AD9364/DMA kabulü", level=2)
        add_code(doc, "cf-ad9361-lpc: Successfully initialized\nad9361_init : AD936x Rev 2 successfully initialized\ncf-ad9361-dds-core-lpc: Successfully initialized\nDMA_EXAMPLE: samples=65536 channels=2 bits=16\nDone.")
        doc.add_heading("8.2 2.4 GHz çevresel aktivite testi", level=2)
        doc.add_paragraph(
            "Hotspot kapalı baseline ve hotspot/traffic açık active taramaları kanallar 1–13 üzerinde "
            "karşılaştırıldı. Her adımda LO, gain, RSSI, bağıl güç, peak, clipping, burst yüzdesi ve DMA "
            "durumu kaydedildi. Örnek kabulde kanal 6, +6.75 dB RSSI, +3.14 dB güç ve +32 yüzde puan "
            "burst farkıyla seçildi; 10 saniyelik monitor sonunda RF_ACTIVITY ve WIFI_RF_TEST PASS verdi."
        )
        doc.add_heading("8.3 Monopulse foundation", level=2)
        add_code(doc, "MONOPULSE_RF: PASS fs=30720000 bw=10000000 lo=2400000000 gain=71 tx_att=89750\nMONOPULSE_FOUNDATION: PASS")
        doc.add_paragraph(
            "Bu işaretçi AD9364 RF profilini ve yazılım altyapısını doğrular; fiziksel dört-durum switch, "
            "faz kalibrasyonu veya açı doğruluğu PASS anlamına gelmez."
        )
        doc.add_heading("8.4 PL FFT kart kabulü", level=2)
        add_code(doc, "PL_FFT: PASS frame=1 bank=0 peak_bin=88 crc=aaa532ad")
        doc.add_paragraph(
            "Sinyalsiz ortamda peak_bin'in frame'ler arasında değişmesi normaldir; maksimum gürültü bini "
            "seçilir. Kontrollü CW kabulünde beklenen davranış, peak'in hesaplanan RF bininde kararlılaşması, "
            "frekans hatasının bin çözünürlüğü içinde kalması ve clipping/CRC hatası oluşmamasıdır."
        )

        add_chapter(doc, 9, "UART protocol v1 ve Python laboratuvar uygulaması")
        add_table(doc, ["Sözleşme", "Değer"], [
            ["Fiziksel hat", "PS_UART, 115200-8-N-1"],
            ["Byte düzeni", "Little-endian; struct padding kullanılmaz"],
            ["Magic", "Komut MPLQ, yanıt MPLR"],
            ["CRC", "CRC-32/ISO-HDLC; poly 0xEDB88320"],
            ["Mesajlar", "GET_INFO, INFO, CAPTURE_ONCE, SPECTRUM, ERROR"],
            ["Spectrum", "Metadata + 2048 adet uint64 magnitude-square"],
            ["Host akışı", "Bir yanıt tamamen doğrulanmadan yeni capture gönderilmez"],
        ], [5.2, 11.6])
        doc.add_heading("9.1 Python katmanları", level=2)
        add_bullets(doc, [
            "Codec: magic/version/type/sequence/payload-length/CRC doğrulaması ve stream resynchronization.",
            "Serial client: timeout, tekrar bağlanma ve bir istek–bir yanıt akışı.",
            "Spectrum model: fftshift, RF frekans ekseni, relative dB, peak doğrulaması ve DC guard.",
            "Viewer: canlı spektrum, son 32 geçerli frame waterfall, metadata ve sağlık göstergeleri.",
            "Persistence: tam frame NPZ; özet ve peak alanları CSV olarak kaydedilir."
        ])
        doc.add_heading("9.2 İzlenen sağlık alanları", level=2)
        add_table(doc, ["Alan", "Yorum"], [
            ["frame_id / sequence", "Frame kaybı veya tekrarı"],
            ["frame_crc32 / message CRC", "PL frame ve taşıma bütünlüğü"],
            ["sample_count / fft_size", "Eksik veya yanlış boyutlu frame"],
            ["skipped_sample_count", "Shadow capture örnek kaybı"],
            ["reset_epoch", "PL reset sonrasında eski/yeni frame ayrımı"],
            ["active_bank / rf_state", "Atomik bank ve gelecekte S0–S3 eşlemesi"],
            ["peak_re / peak_im", "Sayısal teşhis; mutlak faz kalibrasyonu değildir"],
        ], [5.2, 11.6])

        add_chapter(doc, 10, "Sürekli UDP/Ethernet FFT yolu")
        doc.add_picture(str(flow_udp), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_table(doc, ["Ağ/taşıma parametresi", "Değer"], [
            ["Kart IP", "192.168.1.10"], ["PC IP", "192.168.1.100/24"],
            ["UDP port", "5005"], ["Nominal hız", "4 tam frame/s"],
            ["Datagram/frame", "12"], ["Header", "64 byte"],
            ["Frame payload", "2048 × uint64 spectrum + metadata"],
            ["Varsayılan süre", "Sınırsız; --frames ile sınırlandırılabilir"],
        ], [6.0, 10.8])
        doc.add_heading("10.1 Uygulanan yazılım yapısı", level=2)
        add_bullets(doc, [
            "Ayrı vitis_ethernet workspace ve ayrı PL_FFT_ETHERNET compile-time modu oluşturuldu.",
            "GEM3/lwIP211 RAW UDP yolu Golden, vitis2 ve mevcut Candidate paketlerine dokunmadan eklendi.",
            "Taşımadan bağımsız sürekli FFT servisi atomik frame snapshotını packetizer'a verir.",
            "Python decoder datagram başlığı, sıra, parça ofseti, kapsam ve CRC'yi doğrular.",
            "Reassembler eksik, tekrar, sırası değişmiş veya çakışan parçaları teşhis eder; eksik frame'i çizmez.",
            "Viewer süresiz canlı kullanım; headless acceptance sınırlı frame; soak aracı uzun süreli güvenilirlik içindir."
        ])
        doc.add_heading("10.2 Fiziksel kabul prosedürü", level=2)
        add_numbered(doc, [
            "USB-Ethernet adaptörünü 192.168.1.100/24 statik IP'ye ayarla; gateway ve DNS boş bırak.",
            "Kartı JTAG + PS_UART + Ethernet ile bağla; Golden regresyonundan sonra Ethernet Dev launcher'ı çalıştır.",
            "UART READY işaretçisini ve ping 192.168.1.10 sonucunu doğrula.",
            "150 tam frame headless acceptance çalıştır; CRC, sıra, eksik parça ve timeout sıfır olmalı.",
            "En az 600 saniye UDP soak çalıştır; drop/reorder/duplicate istatistiklerini kaydet.",
            "Ağ kabulü geçerse düşük seviyeli, limitli kablosuz SMW200A CW uygula ve beklenen bin/frekans kararlılığını doğrula.",
            "Yalnız bütün kapılar geçerse monopulse-live-ethernet-v1 Candidate paketini oluştur."
        ])
        add_code(doc, ".\\PROJECT\\tools\\Start-Axu4evAd9364EthernetDev.cmd\n.\\PROJECT\\tools\\Start-Axu4evAd9364EthernetViewer.cmd")

        add_chapter(doc, 11, "ZX80/HMC936 ve dört-durum monopulse entegrasyonu")
        doc.add_heading("11.1 Hedef analog durumlar", level=2)
        add_table(doc, ["Durum", "ZX80 seçimi", "Faz kodu", "Ölçülen spektrum"], [
            ["S0", "ΔAZ yolu", "0°", "P_AZ+ = |Σ+ΔAZ|²"],
            ["S1", "ΔAZ yolu", "180°", "P_AZ− = |Σ−ΔAZ|²"],
            ["S2", "ΔEL yolu", "0°", "P_EL+ = |Σ+ΔEL|²"],
            ["S3", "ΔEL yolu", "180°", "P_EL− = |Σ−ΔEL|²"],
        ], [2.5, 4.2, 3.1, 7.2])
        doc.add_heading("11.2 İlk güç tabanlı çıktı", level=2)
        add_code(doc, "e_AZ = (P_AZ+ - P_AZ-) / (P_AZ+ + P_AZ- + epsilon)\ne_EL = (P_EL+ - P_EL-) / (P_EL+ + P_EL- + epsilon)")
        doc.add_paragraph(
            "Dört spectrum aynı ortak CW peak bininde veya onaylanmış dar band-power penceresinde "
            "karşılaştırılır. Bu oranlar Q1.31 olarak PS'ye verilebilir. Açıya dönüşüm yalnız mekanik "
            "-θ/0/+θ kalibrasyonundan sonra LUT/eğri ile yapılır."
        )
        doc.add_heading("11.3 Uygulamadan önce kapatılacak elektriksel kapılar", level=2)
        add_bullets(doc, [
            "ZX80 RF1/RF2/RF COM ve BDC IF kablo takibinin fotoğraf değil süreklilik/etiket kaydıyla kesinleştirilmesi.",
            "ZX80 VDD, Control1/Control2 truth table ve FPGA logic seviyesinin datasheet ile doğrulanması.",
            "HMC936 altı phase bitinin bit ağırlığı/polaritesi, 0° ve 180° kodlarının VNA veya CW ölçümüyle doğrulanması.",
            "5 V level-shifter/buffer, güvenli OE ve power-up default kodunun tasarlanması.",
            "Sekiz GPIO hattının AXU4EVB-P connector/XDC pinlerine bağlanması ve timing-safe PL FSM ile sürülmesi.",
            "Her FFT frame'inin rf_state etiketiyle atomik eşlenmesi; settling/discard sürelerinin ölçülmesi."
        ])
        add_status(doc, "Kritik sınır", "Mevcut live-lab Candidate ZX80 ve HMC936 kontrol pinlerini sürmez. Monopulse foundation PASS, fiziksel S0–S3 veya açı kabulü değildir.", AMBER)

        add_chapter(doc, 12, "Doğrulama stratejisi ve güncel sonuçlar")
        add_table(doc, ["Kapı", "Son taze sonuç", "Kabul kapsamı"], [
            ["Ethernet host testleri", "38 PASS", "Wire codec, reassembly, viewer/acceptance davranışı"],
            ["Monopulse/FFT/UART", "61 PASS + C foundation PASS", "Model, codec, servis ve deterministik C aritmetiği"],
            ["Lab-console", "39 PASS", "Seri akış, model, persistence ve UI yardımcıları"],
            ["Golden release paketi", "PASS", "Manifest/hash ve değişmez release"],
            ["Live-lab Candidate", "PASS", "Paket/launcher bütünlüğü"],
            ["Periodic monitor Candidate", "PASS", "100-frame fiziksel UART kanıtı ve paket"],
            ["Teknik DOCX", "Yapısal + görsel kontrol", "Başlık, tablo, sayfa ve render bütünlüğü"],
        ], [5.1, 4.1, 7.7])
        doc.add_heading("12.1 Bilinen legacy Workspace testi", level=2)
        doc.add_paragraph(
            "Test-Axu4evHandoffArtifacts.ps1 ve Workspace launcher kapısı, farklı tarihlerde üretilmiş "
            "Vivado, Vitis platformu ve IDE bitstream kopyalarının birebir hash eşitliğini arar. Güncel "
            "paralel Candidate tesliminde bu kopyalar aynı jenerasyondan değildir; bu nedenle bu iki "
            "legacy kapı PASS değildir. Günlük teslim bütünlüğü immutable Golden/Candidate paket testleriyle "
            "doğrulanır. Eski Workspace klasörlerinin eşitlenmesi, doğrulanmış paketleri değiştirmeden ayrı bir "
            "bakım işi olarak ele alınmalıdır."
        )
        add_status(doc, "Raporlama ilkesi", "Yazılım/build PASS fiziksel Ethernet veya RF performans PASS yerine kullanılamaz. Açık laboratuvar kapıları açık olarak raporlanmıştır.", LIGHT_BLUE)

        add_chapter(doc, 13, "Tekrarlanabilir çalışma prosedürleri")
        doc.add_heading("13.1 Günlük kart açılışı", level=2)
        add_numbered(doc, [
            "Kart ve AD9364 kartını mekanik olarak kontrol et; JTAG ve PS_UART kablolarını bağla.",
            "PuTTY/terminali 115200-8-N-1 ile aç; yalnız tek hw_server örneği çalıştır.",
            "Golden launcher'ı çalıştır ve DDR + AD9364 başarı işaretçilerini gör.",
            "İstenen deneye göre Candidate launcher'ına geç; Candidate ile bitstream/PMUFW/psu_init/ELF takımını karıştırma.",
            "Tam UART/JTAG kaydını tarihli evidence dosyasına kaydet; yalnız kırpılmış ekran görüntüsüyle kabul yapma.",
            "Test sonunda manifest/hash ve git durumunu kontrol et."
        ])
        doc.add_heading("13.2 Yeni Candidate üretme", level=2)
        add_numbered(doc, [
            "Yeni codex/ önekli dal veya bağımsız worktree oluştur.",
            "Önce sözleşme testi yaz; RED sonucu gözle; minimum uygulamayı ekle; GREEN ve regresyonu çalıştır.",
            "Vivado/Vitis çıktısını ayrı build/workspace'te üret; Golden veya mevcut Candidate'i yerinde değiştirme.",
            "Bitstream, XSA, psu_init, PMUFW ve application ELF kaynak commitini manifestte bağla.",
            "Host testleri ve kart kabul loglarını ayrı kanıt klasörüne ekle.",
            "Yalnız fiziksel kabul geçtiyse board_verified=true kullan; aksi halde false bırak.",
            "Raporları, CURRENT işaretçisini ve teslim notlarını aynı committe güncelle."
        ])
        doc.add_heading("13.3 Teslim bütünlüğü", level=2)
        add_code(doc, "(Get-FileHash .\\FPGA-Profen_AXU4EV_AD9364_Teslim_2026-08-03.rar -Algorithm SHA256).Hash\n# Değer, yanındaki .rar.sha256 dosyasıyla aynı olmalıdır.")

        add_chapter(doc, 14, "Devralma öncelikleri")
        add_table(doc, ["Öncelik", "İş", "Bitti sayılma koşulu"], [
            ["P0", "USB-Ethernet adaptörü ve link", "192.168.1.100/24, UART READY ve ping PASS"],
            ["P0", "UDP fiziksel kabul", "150 frame: sıfır CRC/eksik/sıra hatası"],
            ["P0", "UDP soak", "≥600 s; ölçülmüş kayıp/reorder/duplicate kabul limitlerinde"],
            ["P1", "Kablosuz CW peak", "Beklenen RF bininde kararlı peak; clipping yok"],
            ["P1", "ZX80/HMC936 elektriksel sözleşme", "Kablo matrisi, truth table, level-shifter ve XDC onaylı"],
            ["P2", "S0–S3 PL controller", "Settling/discard + atomik rf_state/frame eşlemesi testli"],
            ["P2", "Normalize AZ/EL", "Dört durumdan ortak-bin Q1.31 oran ve C/Python karşılaştırması"],
            ["P3", "Açı kalibrasyonu", "-θ/0/+θ veri seti, eğri/LUT ve tekrar edilebilir hata bütçesi"],
        ], [2.0, 5.6, 9.2])
        doc.add_heading("14.1 Yapılmaması gerekenler", level=2)
        add_bullets(doc, [
            "PetaLinux/pyadi-iio'yu fiziksel Ethernet kapısını geçmeden zorunlu bağımlılık haline getirmek.",
            "AD9364'ü iki eşzamanlı fiziksel RX kanalı varmış gibi modellemek.",
            "Sinyalsiz peak_bin değişimini FFT arızası olarak sınıflandırmak.",
            "Farklı Candidate'lere ait bitstream, psu_init, PMUFW ve ELF'i aynı başlatmada karıştırmak.",
            "Kalibrasyonsuz bağıl FFT gücünü dBm veya açı doğruluğu olarak sunmak.",
            "Golden veya doğrulanmış Candidate paketlerini yerinde yeniden derlemek."
        ])

        add_chapter(doc, 15, "Dosya ve komut hızlı referansı")
        add_table(doc, ["İhtiyaç", "Dosya / komut"], [
            ["Golden çalıştır", "PROJECT/tools/Start-Axu4evAd9364Golden.cmd"],
            ["UART live-lab", "PROJECT/tools/Start-Axu4evAd9364Candidate.cmd"],
            ["Periyodik 100 frame", "PROJECT/tools/Start-Axu4evAd9364PeriodicMonitor.cmd"],
            ["Ethernet firmware", "PROJECT/tools/Start-Axu4evAd9364EthernetDev.cmd"],
            ["Ethernet viewer", "PROJECT/tools/Start-Axu4evAd9364EthernetViewer.cmd"],
            ["Aktif Vivado", "Related-Docs/hdl-main/projects/fmcomms2/axu4ev/fmcomms2_axu4ev.xpr"],
            ["Nihai durum", "PROJECT/reports/project_situation_report.md"],
            ["Nihai plan", "PROJECT/reports/axu4ev_ad9364_fft_rf_switch_nihai_plan.md"],
            ["Teslim notu", "PROJECT/reports/TESLIM_NOTLARI_2026-08-03.md"],
            ["Genel staj raporu", "PROJECT/reports/profen_staj_raporu.docx"],
            ["Bu teknik rapor", "PROJECT/reports/profen_teknik_detay_raporu.docx"],
        ], [5.0, 11.8])
        doc.add_heading("15.1 Son teknik hüküm", level=2)
        doc.add_paragraph(
            "Teslim edilen platform, AXU4EVB-P + AD9364 üzerinde tekrarlanabilir JTAG bring-up, "
            "2048 nokta PL FFT, peak/spectrum bütünlüğü, UART tabanlı Python gözlem ve ayrı bare-metal "
            "UDP aktarım yazılımını içeren çalışır bir mühendislik temelidir. Teslim, tamamlanmamış "
            "fiziksel Ethernet ve RF switch/phase-shifter ölçümlerini tamamlanmış gibi göstermemekte; "
            "bu kapıları açık, sıralı ve ölçülebilir kabul kriterleriyle devretmektedir."
        )

        # Core properties and update-fields setting.
        props = doc.core_properties
        props.title = "AXU4EVB-P + AD9364 Teknik Tasarım, Uygulama ve Devir Raporu"
        props.subject = "FPGA FFT, bare-metal, UART/UDP ve monopulse teknik devir dokümanı"
        props.author = "FPGA-Profen Proje Ekibi"
        props.keywords = "AXU4EVB-P, AD9364, Vivado, Vitis, FFT, monopulse, UART, UDP"
        settings = doc.settings._element
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)

        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
